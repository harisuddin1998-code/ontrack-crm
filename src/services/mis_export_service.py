# src/services/mis_export_service.py
"""
MIS Export Service - Builds the complete multi-sheet MIS Excel report
(Activities, Installation, Removal, Transfer) and manages its monthly
email delivery.
"""
import os
from datetime import datetime
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd
from flask import current_app

from src.utils.logging import get_logger

logger = get_logger(__name__)

# Column spec = (header, accessor). accessor is an attribute name (str) or a
# callable(obj) -> value for fields that need resolving through a
# relationship (e.g. a linked User's name) rather than a plain attribute.

# Exact column set/order requested for the Activities (REDO) sheet.
ACTIVITIES_SPEC: List[Tuple[str, Any]] = [
    ('ID', 'id'),
    ('ACTIVITY_TYPE', 'activity_type'),
    ('SCHEDULED_DATE', 'scheduled_date'),
    ('RATES', 'rates'),
    ('CUSTOMER_NAME', 'customer_name'),
    ('CUSTOMER_CONTACT', 'customer_contact'),
    ('SALE_PERSON', 'sale_person'),
    ('REGISTRATION_NO', 'registration_no'),
    ('MANUFACTURER', 'make'),
    ('BRAND', 'model'),
    ('YEAR / MODEL', 'year'),
    ('COLOR', 'color'),
    ('CHASSIS_NO', 'chassis_no'),
    ('ENGINE_NO', 'engine_no'),
    ('PREVIOUS_CUSTOMER', 'previous_customer'),
    ('IMEI_NO', 'imei_no'),
    ('SIM_NO', 'sim_no'),
    ('DEVICE_TYPE', 'device_type'),
    ('DEVICE_LOCATION', 'device_location'),
    ('OLD_IMEI_NO', 'old_imei_no'),
    ('OLD_SIM_NO', 'old_sim_no'),
    ('NEW_DEVICE', 'new_device'),
    ('NEW_SIM', 'new_sim'),
    ('T_INSTALLATION', 'transfer_installation'),
    ('TRANSFER_CHARGES', 'transfer_charges'),
    ('OLD_VEHICLE', 'old_vehicle'),
    ('CITY', 'city'),
    ('LOCATION_OF_VEHICLE', 'vehicle_location'),
    ('TECHNICIAN_ASSIGNED', 'technician'),
    ('TESTED_BY', 'tested_by'),
    ('ARRANGED_BY_SALES_PERSON', 'arranged_by'),
    ('FUEL', 'fuel'),
    ('REMARKS', 'remarks'),
    ('RESOLUTION_STATUS', 'resolution_status'),
    ('CREATED_AT', 'created_at'),
]

INSTALLATION_SPEC: List[Tuple[str, Any]] = [
    ('ID', 'id'),
    ('PO_NUMBER', 'po_number'),
    ('ACTIVITY_TYPE', 'activity_type'),
    ('STATUS', 'status'),
    ('OWNER_NAME', 'owner_name'),
    ('OWNER_CONTACT', 'owner_contact'),
    ('CONTACT_PERSON_DRIVER', 'contact_person_driver'),
    ('REGISTRATION_NO', 'reg_no'),
    ('MANUFACTURER', 'vehicle_make'),
    ('BRAND', 'vehicle_model'),
    ('YEAR / MODEL', 'vehicle_year'),
    ('COLOR', 'vehicle_color'),
    ('TRANSMISSION', 'transmission'),
    ('POWER CC', 'power_cc'),
    ('ENGINE_NO', 'engine_number'),
    ('CHASSIS_NO', 'chassis_number'),
    ('SALES_PERSON', lambda po: (po.sales_person.name or po.sales_person.username) if po.sales_person else None),
    ('CITY', 'city'),
    ('VEHICLE_LOCATION', 'vehicle_availability_location'),
    ('EXISTING_CUSTOMER_NAME', 'existing_customer_name'),
    ('EXISTING_VEHICLE_NUMBER', 'existing_vehicle_number'),
    ('SCHEDULED_DATE', 'scheduled_date'),
    ('TECHNICIAN_ASSIGNED', 'technician_assigned'),
    ('IMEI_NO', 'imei_no'),
    ('SIM_NO', 'sim_no'),
    ('DEVICE_TYPE', 'device_type'),
    ('DEVICE_LOCATION', 'device_location'),
    ('FUEL', 'fuel'),
    ('TESTED_BY', 'tested_by'),
    ('ARRANGED_BY_SALES_PERSON', 'arranged_by_sales_person'),
    ('REMARKS', 'remarks'),
    ('RATES', 'rates'),
    ('AMC', 'amc'),
    ('TOTAL_AMOUNT', lambda po: po.get_total_amount()),
    ('COMPLETION_EMAIL_SENT_AT', 'completion_email_sent_at'),
    ('CREATED_AT', 'created_at'),
]

REMOVAL_SPEC: List[Tuple[str, Any]] = [
    ('ID', 'id'),
    ('REGISTRATION_NO', 'registration_no'),
    ('MANUFACTURER', 'make'),
    ('BRAND', 'model'),
    ('YEAR / MODEL', 'year'),
    ('COLOR', 'color'),
    ('CHASSIS_NO', 'chassis_no'),
    ('ENGINE_NO', 'engine_no'),
    ('CUSTOMER_NAME', 'customer_name'),
    ('CUSTOMER_CONTACT', 'customer_contact'),
    ('IMEI_NO', 'imei_no'),
    ('SIM_NO', 'sim_no'),
    ('DEVICE_LOCATION', 'device_location'),
    ('REMOVAL_REASON', 'removal_reason'),
    ('REMOVAL_DATE', 'removal_date'),
    ('REMOVAL_TYPE', 'removal_type'),
    ('DEVICE_RETURNED', 'device_returned'),
    ('RETURN_DATE', 'return_date'),
    ('DEVICE_CONDITION', 'device_condition'),
    ('RETAINED_BY', 'retained_by'),
    ('STORAGE_LOCATION', 'storage_location'),
    ('ASSIGNED_TO', lambda r: r.assigned_to_user.name if r.assigned_to_user else None),
    ('ASSIGNED_BY', lambda r: r.assigned_by_user.name if r.assigned_by_user else None),
    ('ASSIGNED_AT', 'assigned_at'),
    ('STATUS', 'status'),
    ('COMPLETION_NOTES', 'completion_notes'),
    ('COMPLETED_BY', lambda r: r.completed_by_user.name if r.completed_by_user else None),
    ('COMPLETED_AT', 'completed_at'),
    ('CREATED_AT', 'created_at'),
]

DEVICE_CHANGES_SPEC: List[Tuple[str, Any]] = [
    ('ID', 'id'),
    ('REDO_NUMBER', lambda a: a.redo_number or f"REDO-{a.id}"),
    ('DATE', 'created_at'),
    ('REGISTRATION_NO', 'registration_no'),
    ('MANUFACTURER', 'make'),
    ('BRAND', 'model'),
    ('YEAR / MODEL', 'year'),
    ('COLOR', 'color'),
    ('CUSTOMER_NAME', 'customer_name'),
    ('CUSTOMER_CONTACT', 'customer_contact'),
    ('CITY', 'city'),
    ('LOCATION_OF_VEHICLE', 'vehicle_location'),
    ('OLD_IMEI_NO', 'old_imei_no'),
    ('OLD_SIM_NO', 'old_sim_no'),
    ('NEW_DEVICE', 'new_device'),
    ('NEW_SIM', 'new_sim'),
    ('DEVICE_CHANGE_REASON', lambda a: a.device_change_reason or 'Not Specified'),
    ('DEVICE_TYPE', 'device_type'),
    ('DEVICE_LOCATION', 'device_location'),
    ('TECHNICIAN_ASSIGNED', 'technician'),
    ('TESTED_BY', 'tested_by'),
    ('RESOLUTION_STATUS', 'resolution_status'),
    ('REMARKS', 'remarks'),
]

INSTALLATION_RECOVERY_SPEC: List[Tuple[str, Any]] = [
    ('REG_NO', lambda c: c.redo_activity.registration_no if c.redo_activity else None),
    ('CUSTOMER_NAME', lambda c: c.redo_activity.customer_name if c.redo_activity else None),
    ('REASON', 'reason'),
    ('AMOUNT_DUE', 'amount'),
    ('STATUS', 'status'),
    ('ASSIGNED_OFFICER', lambda c: c.assigned_officer.name if c.assigned_officer else None),
    ('PAYMENT_METHOD', 'payment_method'),
    ('PAYMENT_REFERENCE', 'payment_reference'),
    ('RECOVERED_AT', 'recovered_at'),
    ('CREATED_AT', 'created_at'),
]

INSTALLATION_RECOVERY_FOLLOWUP_SPEC: List[Tuple[str, Any]] = [
    ('DATE', 'conversation_date'),
    ('REG_NO', lambda f: f.charge.redo_activity.registration_no if f.charge and f.charge.redo_activity else None),
    ('CUSTOMER_NAME', lambda f: f.charge.redo_activity.customer_name if f.charge and f.charge.redo_activity else None),
    ('REASON', lambda f: f.charge.reason if f.charge else None),
    ('AMOUNT_DUE', lambda f: f.charge.amount if f.charge else None),
    ('CHARGE_STATUS', lambda f: f.charge.status if f.charge else None),
    ('CHANNEL', 'conversation_type'),
    ('DIRECTION', 'direction'),
    ('CONTACT_PERSON', 'contact_person'),
    ('CONTACT_NUMBER', 'contact_number'),
    ('SUMMARY', 'summary'),
    ('ACTION_TAKEN', 'action_taken'),
    ('FOLLOW_UP_DATE', 'follow_up_date'),
    ('RECORDED_BY', 'recorded_by_name'),
]

def _scheduled_redo_for(conversation):
    """The REDO scheduled off the back of a follow-up conversation, if any.

    Non-reporting follow-ups only become a formal REDO when a technician is
    actually assigned (see redo.non_reporting_detail), and the link between
    the two is the registration number rather than a foreign key. The most
    recent REDO for that vehicle is the one the follow-up produced.
    """
    from src.models.redo import RedoActivity
    vehicle = getattr(conversation, 'vehicle', None)
    reg_no = getattr(vehicle, 'registration_no', None)
    if not reg_no:
        return None
    return (RedoActivity.query
            .filter(RedoActivity.registration_no.ilike(f"%{reg_no}%"))
            .order_by(RedoActivity.created_at.desc())
            .first())


def _redo_scheduled_by(conversation):
    """Name of whoever scheduled the REDO for this vehicle.

    `assigned_by` is set to the signed-in user at the moment the REDO is
    raised, so it is the scheduler - not the technician it was handed to.
    """
    redo = _scheduled_redo_for(conversation)
    if redo is None:
        return None
    if redo.assigned_by_user:
        return redo.assigned_by_user.name or redo.assigned_by_user.username
    return redo.arranged_by


FOLLOWUP_SPEC: List[Tuple[str, Any]] = [
    ('DATE', 'conversation_date'),
    ('REGISTRATION_NO', lambda c: c.vehicle.registration_no if c.vehicle else None),
    ('CUSTOMER_NAME', lambda c: c.vehicle.customer_name if c.vehicle else None),
    ('CUSTOMER_CONTACT', lambda c: c.vehicle.customer_contact if c.vehicle else None),
    ('CITY', lambda c: c.vehicle.city if c.vehicle else None),
    ('MANUFACTURER', lambda c: c.vehicle.make if c.vehicle else None),
    ('BRAND', lambda c: c.vehicle.model if c.vehicle else None),
    ('DAYS_NON_REPORTING', lambda c: c.vehicle.get_days_non_reporting() if c.vehicle else None),
    ('CONTACT_OUTCOME', lambda c: c.vehicle.contact_outcome if c.vehicle else None),
    ('CHANNEL', 'conversation_type'),
    ('DIRECTION', 'direction'),
    ('CONTACT_PERSON', 'contact_person'),
    ('CONTACT_NUMBER', 'contact_number'),
    ('SUMMARY', 'summary'),
    ('ACTION_TAKEN', 'action_taken'),
    ('FOLLOW_UP_DATE', 'follow_up_date'),
    ('REDO_NUMBER', lambda c: (lambda r: r.redo_number if r else None)(_scheduled_redo_for(c))),
    ('REDO_SCHEDULED_BY', _redo_scheduled_by),
    ('REDO_TECHNICIAN', lambda c: (lambda r: r.technician if r else None)(_scheduled_redo_for(c))),
    ('LOGGED_BY', 'recorded_by_name'),
]

TRANSFER_SPEC: List[Tuple[str, Any]] = [
    ('ID', 'id'),
    ('OLD_REGISTRATION_NO', 'old_registration_no'),
    ('OLD_MANUFACTURER', 'old_make'),
    ('OLD_BRAND', 'old_model'),
    ('OLD_YEAR_MODEL', 'old_year'),
    ('OLD_COLOR', 'old_color'),
    ('OLD_CHASSIS_NO', 'old_chassis_no'),
    ('OLD_ENGINE_NO', 'old_engine_no'),
    ('OLD_CUSTOMER_NAME', 'old_customer_name'),
    ('OLD_CUSTOMER_CONTACT', 'old_customer_contact'),
    ('OLD_IMEI_NO', 'old_imei_no'),
    ('OLD_SIM_NO', 'old_sim_no'),
    ('OLD_DEVICE_LOCATION', 'old_device_location'),
    ('NEW_REGISTRATION_NO', 'new_registration_no'),
    ('NEW_MANUFACTURER', 'new_make'),
    ('NEW_BRAND', 'new_model'),
    ('NEW_YEAR_MODEL', 'new_year'),
    ('NEW_COLOR', 'new_color'),
    ('NEW_CHASSIS_NO', 'new_chassis_no'),
    ('NEW_ENGINE_NO', 'new_engine_no'),
    ('NEW_CUSTOMER_NAME', 'new_customer_name'),
    ('NEW_CUSTOMER_CONTACT', 'new_customer_contact'),
    ('TRANSFER_REASON', 'transfer_reason'),
    ('DEVICE_TRANSFERRED_DATE', 'device_transferred_date'),
    ('ASSIGNED_TO', lambda r: r.assigned_to_user.name if r.assigned_to_user else None),
    ('ASSIGNED_BY', lambda r: r.assigned_by_user.name if r.assigned_by_user else None),
    ('ASSIGNED_AT', 'assigned_at'),
    ('STATUS', 'status'),
    ('COMPLETION_NOTES', 'completion_notes'),
    ('COMPLETED_BY', lambda r: r.completed_by_user.name if r.completed_by_user else None),
    ('COMPLETED_AT', 'completed_at'),
    ('CREATED_AT', 'created_at'),
]


def _extract_rows(objects: List[Any], spec: List[Tuple[str, Any]]) -> pd.DataFrame:
    """Build a DataFrame from a list of ORM objects per a (header, accessor)
    spec. Headers always appear even when `objects` is empty, since the
    column list comes from the spec, not from inspecting the data."""
    headers = [header for header, _ in spec]
    rows: List[Dict[str, Any]] = []
    for obj in objects:
        row = {}
        for header, accessor in spec:
            row[header] = accessor(obj) if callable(accessor) else getattr(obj, accessor, None)
        rows.append(row)
    return pd.DataFrame(rows, columns=headers)


class MISExportService:
    """Builds the complete MIS Excel workbook (Activities, Installation,
    Removal, Transfer) and manages its monthly email delivery."""

    def __init__(self):
        upload_setting = current_app.config.get('UPLOAD_FOLDER', 'uploads')
        if not os.path.isabs(upload_setting):
            project_root = os.path.abspath(os.path.join(current_app.root_path, '..'))
            self.upload_folder = os.path.abspath(os.path.join(project_root, upload_setting))
        else:
            self.upload_folder = os.path.abspath(upload_setting)
        self.export_folder = os.path.join(self.upload_folder, 'mis_reports')
        os.makedirs(self.export_folder, exist_ok=True)

    def generate_mis_excel(self, date_range: Optional[Dict[str, Any]] = None) -> str:
        """Build the multi-sheet MIS workbook and return its filepath.

        `date_range` is a parsed range from `src.utils.date_ranges` - pass one
        to cover a specific FROM/TO period (what the download picker and the
        monthly email both do), or omit it for everything ever recorded.
        Every sheet is filtered on its own `created_at`, so one period means
        the same period on every sheet.
        """
        from src.models.redo import RedoActivity
        from src.models.purchase_order import PurchaseOrder
        from src.models.removal import RemovalRetainedActivity, RemovalTransferActivity
        from src.models.installation_recovery import InstallationRecoveryCharge
        from src.services.data_validation_service import DataValidationService
        from src.utils.date_ranges import apply_range

        date_range = date_range or {}
        suffix = ''
        if date_range.get('from_str') and date_range.get('to_str'):
            suffix = f"_{date_range['from_str']}_to_{date_range['to_str']}"
        filename = f"MIS_Report{suffix}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.export_folder, filename)

        def scoped(model):
            return apply_range(model.query, model.created_at, date_range).order_by(model.id)

        # A backend validation pass gates the official MIS totals - no human
        # click-to-verify step. Each REDO activity is checked automatically
        # (required fields present, completion details consistent); records
        # that fail stay excluded from Activities/Device Changes and are
        # listed with their specific reason on a "Data Issues" sheet.
        all_activities = scoped(RedoActivity).all()
        activities, invalid_activities = DataValidationService().split_redo_activities(all_activities)

        installations = scoped(PurchaseOrder).all()
        removals = scoped(RemovalRetainedActivity).all()
        transfers = scoped(RemovalTransferActivity).all()
        device_changes = [
            a for a in activities
            if (a.new_device and a.new_device.strip()) or (a.new_sim and a.new_sim.strip())
        ]
        installation_recovery_charges = scoped(InstallationRecoveryCharge).all()

        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            _extract_rows(activities, ACTIVITIES_SPEC).to_excel(writer, sheet_name='Activities', index=False)
            _extract_rows(installations, INSTALLATION_SPEC).to_excel(writer, sheet_name='Installation', index=False)
            _extract_rows(removals, REMOVAL_SPEC).to_excel(writer, sheet_name='Removal', index=False)
            _extract_rows(transfers, TRANSFER_SPEC).to_excel(writer, sheet_name='Transfer', index=False)
            _extract_rows(device_changes, DEVICE_CHANGES_SPEC).to_excel(writer, sheet_name='Device Changes', index=False)
            _extract_rows(installation_recovery_charges, INSTALLATION_RECOVERY_SPEC).to_excel(
                writer, sheet_name='Installation Recovery', index=False)

            issues_rows = [
                {'REDO_ID': a.id, 'REGISTRATION_NO': a.registration_no, 'ISSUES': '; '.join(errs)}
                for a, errs in invalid_activities
            ]
            pd.DataFrame(issues_rows, columns=['REDO_ID', 'REGISTRATION_NO', 'ISSUES']).to_excel(
                writer, sheet_name='Data Issues', index=False)

        self._autofit_columns(filepath)
        logger.info(
            f"MIS Excel report generated: {filepath} "
            f"(period={date_range.get('label', 'All time')}, "
            f"Activities={len(activities)}, Installation={len(installations)}, "
            f"Removal={len(removals)}, Transfer={len(transfers)}, "
            f"DeviceChanges={len(device_changes)}, "
            f"data_issues_excluded={len(invalid_activities)})"
        )
        return filepath

    # ------------------------------------------------------------------
    # Standalone restricted reports
    # ------------------------------------------------------------------

    def generate_device_change_excel(self, date_range: Optional[Dict[str, Any]] = None,
                                     changes: Optional[List[Any]] = None) -> str:
        """Device Change Report as its own workbook.

        Also a sheet inside the full MIS workbook; this is the standalone
        download and the file attached to the monthly management email.
        Pass `changes` to export exactly what a screen is showing (filters
        and all), or a `date_range` to let the service select the period.
        """
        from src.models.redo import RedoActivity
        from src.utils.date_ranges import apply_range

        if changes is None:
            is_device_change = RedoActivity.query.filter(
                (RedoActivity.new_device.isnot(None) & (RedoActivity.new_device != ''))
                | (RedoActivity.new_sim.isnot(None) & (RedoActivity.new_sim != ''))
            )
            changes = apply_range(is_device_change, RedoActivity.created_at, date_range or {}) \
                .order_by(RedoActivity.created_at.desc()).all()

        filename = f"Device_Change_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.export_folder, filename)
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            _extract_rows(changes, DEVICE_CHANGES_SPEC).to_excel(
                writer, sheet_name='Device Changes', index=False)
        self._autofit_columns(filepath)
        return filepath

    def generate_technician_activity_excel(self, rows: List[Dict[str, Any]]) -> str:
        """Technician Activity Report from the wallboard's own rows.

        Takes the computed rows rather than re-querying, so the file and the
        wallboard can never disagree about a technician's numbers.
        """
        filename = f"Technician_Activity_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.export_folder, filename)
        frame = pd.DataFrame(rows)
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            frame.to_excel(writer, sheet_name='Technician Activity', index=False)
        self._autofit_columns(filepath)
        return filepath

    def generate_followup_excel(self, conversations: List[Any]) -> str:
        """REDO Follow-up Report download.

        Includes who scheduled the REDO for each vehicle - the report is used
        to chase up follow-ups, and the first question is always whose it was.
        """
        filename = f"REDO_Followup_Report_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.export_folder, filename)
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            _extract_rows(conversations, FOLLOWUP_SPEC).to_excel(
                writer, sheet_name='Follow-up Report', index=False)
        self._autofit_columns(filepath)
        return filepath

    def generate_installation_recovery_excel(self, charges: List[Any]) -> str:
        """Standalone single-sheet export for the Installation Recovery
        report page's own download button, separate from the full MIS
        workbook above (which also includes this data as one of its sheets)."""
        filename = f"Installation_Recovery_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.export_folder, filename)
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            _extract_rows(charges, INSTALLATION_RECOVERY_SPEC).to_excel(
                writer, sheet_name='Installation Recovery', index=False)
        self._autofit_columns(filepath)
        return filepath

    def generate_installation_recovery_followup_excel(self, followups: List[Any]) -> str:
        """Installation Recovery Follow-up Report download - the same
        pattern as the REDO Follow-up Report, but for device recovery calls."""
        filename = f"Installation_Recovery_Followup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.export_folder, filename)
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            _extract_rows(followups, INSTALLATION_RECOVERY_FOLLOWUP_SPEC).to_excel(
                writer, sheet_name='Follow-up Report', index=False)
        self._autofit_columns(filepath)
        return filepath

    @staticmethod
    def _autofit_columns(filepath: str) -> None:
        """Make a generated workbook presentable before it leaves the server.

        Every export in this service ends here, so the whole family looks the
        same: a bold, centred, filled header row that stays visible while you
        scroll, filters on it, columns wide enough to read, and each column
        aligned to what it holds - numbers right so the digits line up, dates
        centred, text left. A spreadsheet that is mailed to senior management
        on the 1st should not need reformatting before it can be read.
        """
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter

        HEADER_FILL = PatternFill('solid', fgColor='5478C0')   # the CRM's list header
        HEADER_FONT = Font(bold=True, color='FFFFFF', size=11)
        HEADER_ALIGN = Alignment(horizontal='center', vertical='center', wrap_text=True)
        THIN = Side(style='thin', color='D6DEEF')
        BODY_BORDER = Border(bottom=THIN)

        wb = load_workbook(filepath)
        for ws in wb.worksheets:
            if ws.max_row < 1:
                continue

            for cell in ws[1]:
                cell.fill = HEADER_FILL
                cell.font = HEADER_FONT
                cell.alignment = HEADER_ALIGN
            ws.row_dimensions[1].height = 30

            # Freeze under the header and filter on it - on a sheet of a few
            # thousand rows both are the difference between a report and a
            # dump.
            ws.freeze_panes = 'A2'
            if ws.max_column >= 1:
                # Over the header even on an empty sheet - a report with no
                # rows this month still opens with the same controls as one
                # that has them.
                ws.auto_filter.ref = (f'A1:{get_column_letter(ws.max_column)}'
                                      f'{max(ws.max_row, 1)}')

            for col_idx, col_cells in enumerate(ws.columns, start=1):
                body = list(col_cells)[1:]
                width = max((len(str(c.value)) if c.value is not None else 0)
                            for c in col_cells)
                ws.column_dimensions[get_column_letter(col_idx)].width = \
                    min(max(width + 3, 11), 42)

                # Align on what the column actually holds rather than on its
                # name - a "RATES" column of blanks should not be right
                # aligned just because of what it is called.
                values = [c.value for c in body if c.value not in (None, '')]
                numeric = values and all(isinstance(v, (int, float)) and
                                         not isinstance(v, bool) for v in values)
                dated = values and all(hasattr(v, 'strftime') for v in values)
                horizontal = 'right' if numeric else ('center' if dated else 'left')

                for cell in body:
                    cell.alignment = Alignment(horizontal=horizontal, vertical='center')
                    cell.border = BODY_BORDER
                    if numeric and isinstance(cell.value, float):
                        cell.number_format = '#,##0.00'
                    elif numeric:
                        cell.number_format = '#,##0'
                    elif dated:
                        cell.number_format = 'dd mmm yyyy'

        wb.save(filepath)

    def generate_report_excel(self, title: str, columns: List[Dict[str, Any]],
                              rows: List[Dict[str, Any]]) -> str:
        """Download for a catalogue report, from the same columns and rows the
        screen rendered - so the file and the page can never disagree about
        what the report contains or what order it is in."""
        safe = ''.join(ch if ch.isalnum() else '_' for ch in title)[:40].strip('_')
        filename = f"{safe}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.export_folder, filename)

        headers = [col['label'] for col in columns]
        data = [[row.get(col['key']) for col in columns] for row in rows]

        # Excel sheet names cannot exceed 31 characters or contain []:*?/\
        sheet = self._sheet_name(title)
        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            pd.DataFrame(data, columns=headers).to_excel(
                writer, sheet_name=sheet, index=False)
        self._autofit_columns(filepath)
        return filepath

    def generate_multi_sheet_report_excel(self, title: str,
                                          tables: List[Dict[str, Any]]) -> str:
        """One sheet per table, for the reports that run to several - the
        monthly workbook above all. Same styling as every other export."""
        safe = ''.join(ch if ch.isalnum() else '_' for ch in title)[:40].strip('_')
        filename = f"{safe}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        filepath = os.path.join(self.export_folder, filename)

        with pd.ExcelWriter(filepath, engine='openpyxl') as writer:
            for index, table in enumerate(tables):
                columns = table['columns']
                frame = pd.DataFrame(
                    [[row.get(col['key']) for col in columns] for row in table['rows']],
                    columns=[col['label'] for col in columns])
                frame.to_excel(writer,
                               sheet_name=self._sheet_name(table.get('title')
                                                           or f'Sheet {index + 1}'),
                               index=False)
        self._autofit_columns(filepath)
        return filepath

    @staticmethod
    def _sheet_name(title: str) -> str:
        """Excel rejects sheet names over 31 characters or containing []:*?/\\."""
        cleaned = ''.join(ch for ch in title if ch not in '[]:*?/\\')
        return cleaned[:31].strip() or 'Report'

    # ------------------------------------------------------------------
    # Monthly management distribution
    #
    # All three reports below are restricted to senior management - see
    # SENIOR_MANAGEMENT_RECIPIENTS in email_service. They run on the 1st and
    # cover the month that just finished, not the one that started this
    # morning: on 1 September the report is August's.
    # ------------------------------------------------------------------

    XLSX_MIME = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'

    @staticmethod
    def _previous_month() -> Dict[str, Any]:
        from src.utils.date_ranges import parse_range, previous_month_bounds
        start, end = previous_month_bounds()
        return parse_range({'from': start.strftime('%Y-%m-%d'), 'to': end.strftime('%Y-%m-%d')})

    def _attachment(self, filepath: str) -> Dict[str, str]:
        return {
            'filename': os.path.basename(filepath),
            'filepath': filepath,
            'content_type': self.XLSX_MIME,
        }

    def email_monthly_report(self) -> bool:
        """Email last month's complete MIS report to senior management."""
        from src.models.redo import RedoActivity
        from src.services.email_service import EmailService
        from src.services.data_validation_service import DataValidationService
        from src.utils.date_ranges import apply_range

        period = self._previous_month()
        month_label = period['start'].strftime('%B %Y')
        filepath = self.generate_mis_excel(period)

        month_activities = apply_range(RedoActivity.query, RedoActivity.created_at, period).all()
        _valid, invalid_activities = DataValidationService().split_redo_activities(month_activities)
        data_issues_count = len(invalid_activities)

        pending_note = (
            f"<p style='color:#c2410c;'><b>Note:</b> {data_issues_count} REDO "
            f"activity(ies) failed automatic data validation and were excluded from "
            f"this report's Activities/Device Changes totals - see the 'Data Issues' sheet.</p>"
            if data_issues_count else ""
        )
        return EmailService().send_restricted_report(
            subject=f"Monthly MIS Report - {month_label}",
            html_body=(
                f"<p>Please find attached the complete MIS report for <b>{month_label}</b> "
                f"({period['label']}), covering Activities, Installation, Removal, and "
                f"Transfer records.</p>"
                f"{pending_note}"
                f"<p>This is an automated report generated on the 1st of the month.</p>"
            ),
            attachments=[self._attachment(filepath)],
        )

    def email_monthly_device_change_report(self) -> bool:
        """Email last month's Device Change report to senior management."""
        from src.services.email_service import EmailService

        period = self._previous_month()
        month_label = period['start'].strftime('%B %Y')
        filepath = self.generate_device_change_excel(period)

        return EmailService().send_restricted_report(
            subject=f"Device Change Report - {month_label}",
            html_body=(
                f"<p>Please find attached the Device Change report for <b>{month_label}</b> "
                f"({period['label']}) - every REDO in which a device or SIM was replaced, "
                f"with the recorded reason.</p>"
                f"<p>This is an automated report generated on the 1st of the month.</p>"
            ),
            attachments=[self._attachment(filepath)],
        )

    def email_monthly_report_pack(self) -> bool:
        """Email every report in the catalogue, for the month just ended.

        One job and one email rather than one of each per report: the reports
        are read together at month end, and a single send cannot half-fail
        and leave management holding three of nine.

        The catalogue is the source of what goes out, so a report added to
        the CRM is in the month-end pack without anyone remembering to add
        it here.
        """
        from src.services.email_service import EmailService
        from src.web.admin import (MIS_REPORTS, MIS_REPORT_BUILDERS,
                                   MIS_REPORTS_BY_SLUG, _report_tables)

        period = self._previous_month()
        month_label = period['start'].strftime('%B %Y')

        attachments: List[Dict[str, str]] = []
        included: List[str] = []
        failed: List[str] = []

        def attach(name: str, build) -> None:
            try:
                attachments.append(self._attachment(build()))
                included.append(name)
            except Exception as exc:            # noqa: BLE001 - one bad report
                # must not stop the other eight going out.
                logger.error('Month-end pack: %s failed to build: %s', name, exc)
                failed.append(name)

        # Monthly MIS is attached by hand because it is a multi-sheet workbook
        # assembled its own way, not from a catalogue builder's tables.
        #
        # Device Changes, Technician Performance and REDO Follow-Ups used to be
        # attached here too, since they had no builders. They have builders
        # now, so the loop below covers them - attaching them here as well sent
        # senior management two copies of each under the same name.
        attach(MIS_REPORTS_BY_SLUG['monthly-mis']['name'],
               lambda: self.generate_mis_excel(period))

        for report in MIS_REPORTS:
            builder = MIS_REPORT_BUILDERS.get(report['slug'])
            if builder is None or report['slug'] == 'monthly-mis':
                continue        # no builder, or already attached above

            def build(_report=report, _builder=builder):
                tables = _report_tables(_builder(period))
                if len(tables) == 1:
                    return self.generate_report_excel(
                        _report['name'], tables[0]['columns'], tables[0]['rows'])
                return self.generate_multi_sheet_report_excel(_report['name'], tables)

            attach(report['name'], build)

        listed = ''.join(f'<li>{name}</li>' for name in included)
        problem = (f"<p style='color:#c2410c;'><b>Note:</b> could not be built this "
                   f"month: {', '.join(failed)}.</p>" if failed else '')

        return EmailService().send_restricted_report(
            subject=f"Month-End Report Pack - {month_label}",
            html_body=(
                f"<p>Attached are the month-end reports for <b>{month_label}</b> "
                f"({period['label']}).</p>"
                f"<ul>{listed}</ul>"
                f"{problem}"
                f"<p>Generated automatically on the 1st of the month at 00:00.</p>"
            ),
            attachments=attachments,
        )

    def email_monthly_technician_activity_report(self) -> bool:
        """Email last month's Technician Activity report to senior management."""
        from src.services.email_service import EmailService
        from src.services.technician_activity_service import TechnicianActivityService

        period = self._previous_month()
        month_label = period['start'].strftime('%B %Y')
        rows = TechnicianActivityService().wallboard_rows(period)
        filepath = self.generate_technician_activity_excel(rows)

        return EmailService().send_restricted_report(
            subject=f"Technician Activity Report - {month_label}",
            html_body=(
                f"<p>Please find attached the Technician Activity report for "
                f"<b>{month_label}</b> ({period['label']}) - installations, REDOs, "
                f"removals and transfers per technician, across every status.</p>"
                f"<p>This is an automated report generated on the 1st of the month.</p>"
            ),
            attachments=[self._attachment(filepath)],
        )
